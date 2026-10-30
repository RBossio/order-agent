"""Generatore 150 ordini con documenti PDF, XLSX, EML e ground truth."""

import hashlib
import json
import random
import re
import uuid
from datetime import datetime, timedelta
from email.mime.application import MIMEApplication
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from io import BytesIO
from pathlib import Path

from . import SEED

NAMESPACE_VENTURA = uuid.UUID("a1b2c3d4-e5f6-7890-abcd-ef1234567890")

FIXED_GENERATION_TIMESTAMP = datetime(2026, 1, 15, 10, 30, 0)


def _deterministic_uuid(order_idx: int, suffix: str = "") -> str:
    """Genera un UUID deterministico basato su order_idx e suffix."""
    name = f"order-{order_idx}-{suffix}-seed{SEED}"
    return str(uuid.uuid5(NAMESPACE_VENTURA, name))


def _deterministic_doc_id(order_idx: int, doc_type: str) -> str:
    """Genera un document_id deterministico."""
    u = uuid.uuid5(NAMESPACE_VENTURA, f"doc-{order_idx}-{doc_type}-seed{SEED}")
    return f"doc-{u.hex[:8]}"


def _fixed_timestamp_iso() -> str:
    """Restituisce il timestamp fisso in formato ISO."""
    return FIXED_GENERATION_TIMESTAMP.strftime("%Y-%m-%dT%H:%M:%SZ")


def _fixed_email_date() -> str:
    """Restituisce il timestamp fisso per header email."""
    return FIXED_GENERATION_TIMESTAMP.strftime("%a, %d %b %Y %H:%M:%S +0000")

NUM_IN_WORDS = {
    1: "one", 2: "two", 3: "three", 4: "four", 5: "five",
    6: "six", 7: "seven", 8: "eight", 9: "nine", 10: "ten",
    11: "eleven", 12: "twelve", 15: "fifteen", 20: "twenty",
    24: "twenty-four", 25: "twenty-five", 50: "fifty", 100: "one hundred"
}

HARD_CASE_TYPES = [
    "UOM_CONVERSION",
    "UOM_CARTON",
    "CUSTOMER_CODES",
    "UNKNOWN_CUSTOMER_CODE",
    "AMBIGUOUS_CUSTOMER_CODE",
    "PRICE_OUT_OF_TOLERANCE",
    "DUP_PO",
    "AMBIGUOUS_DESC",
    "TYPO_DESC",
    "TOTAL_MIXED",
    "MULTIPAGE_PDF",
    "DATE_FORMAT",
    "QTY_WORDS",
    "EMAIL_CONTRADICTORY",
    "MISSING_PO_NUMBER",
    "MISSING_QTY",
    "MISSING_DATE",
]


def assign_hard_cases(
    num_orders: int,
    hard_case_pct: float,
    types: list[str] | None = None,
    rng: random.Random | None = None,
) -> list[list[str]]:
    """Assegna i tipi di caso difficile a un sottoinsieme di ordini.

    Ogni tipo compare almeno una volta se ``num_hard >= len(types)``.
    """
    pool = list(types if types is not None else HARD_CASE_TYPES)
    picker = rng if rng is not None else random
    num_hard = int(num_orders * hard_case_pct)
    num_normal = num_orders - num_hard
    assignments: list[list[str]] = []
    for hc_type in pool:
        assignments.append([hc_type])
    while len(assignments) < num_hard and pool:
        combo_size = picker.randint(1, min(2, len(pool)))
        combo = picker.sample(pool, combo_size)
        assignments.append(combo)
    picker.shuffle(assignments)
    assignments = assignments[:num_hard]
    configs: list[list[str]] = [[] for _ in range(num_normal)]
    configs.extend(assignments)
    picker.shuffle(configs)
    return configs


def carton_conversion(item: dict) -> dict | None:
    """Conversione da cartone (CS) coerente con ``products.json``.

    Preferisce CS → BX, poi le altre unità interne usate dai dati.
    """
    conversions = item.get("conversions") or []
    cs_convs = [
        conv
        for conv in conversions
        if str(conv.get("from_uom") or "").upper() == "CS" and conv.get("factor")
    ]
    for target in ("BX", "EA", "RL", "RO", "GLL", "PK"):
        for conv in cs_convs:
            if str(conv.get("to_uom") or "").upper() == target:
                return {
                    "from_uom": "CS",
                    "to_uom": conv["to_uom"],
                    "factor": float(conv["factor"]),
                }
    if cs_convs:
        conv = cs_convs[0]
        return {
            "from_uom": "CS",
            "to_uom": conv["to_uom"],
            "factor": float(conv["factor"]),
        }
    pack = item.get("pack") or {}
    if str(item.get("uom") or "").upper() == "CS" and pack.get("inner_count"):
        return {
            "from_uom": "CS",
            "to_uom": pack.get("inner_uom") or "EA",
            "factor": float(pack["inner_count"]),
        }
    return None


def _sha256(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


_DATE_PARSE_FORMATS = (
    "%Y-%m-%d",
    "%m/%d/%Y",
    "%d-%b-%Y",
    "%m-%d-%Y",
    "%d/%m/%Y",
    "%d.%m.%Y",
    "%b %d, %Y",
    "%Y/%m/%d",
    "%d-%b-%y",
    "%m/%d/%y",
    "%B %d, %Y",
    "%Y.%m.%d",
)

EMAIL_BODY_TEMPLATES = (
    "free_text",
    "html_table",
    "signature_list",
    "qty_words",
    "forwarded_reply",
)


def _parse_display_date(text: str) -> datetime | None:
    """Parsa una data già formattata dal generatore."""
    raw = (text or "").strip()
    if not raw:
        return None
    for fmt in _DATE_PARSE_FORMATS:
        try:
            return datetime.strptime(raw, fmt)
        except ValueError:
            continue
    return None


def _ordinal_day(day: int) -> str:
    if 10 <= day % 100 <= 20:
        suffix = "th"
    else:
        suffix = {1: "st", 2: "nd", 3: "rd"}.get(day % 10, "th")
    return f"{day}{suffix}"


def reformat_existing_date(date_display: str, kind: str) -> str:
    """Riformatta la data reale (DATE_FORMAT), senza sostituirla con un valore a caso."""
    parsed = _parse_display_date(date_display)
    if parsed is None:
        return date_display
    variants = {
        "pdf": ["%b %d, %Y", "%d-%b-%y", "%m/%d/%y", "%d.%m.%Y"],
        "xlsx": ["%d-%b-%Y", "%m/%d/%y", "%B %d, %Y"],
        "eml": ["ordinal", "%d/%m/%Y", "%Y.%m.%d"],
    }
    fmt = random.choice(variants.get(kind, ["%Y-%m-%d"]))
    if fmt == "ordinal":
        return f"{parsed.strftime('%b')} {_ordinal_day(parsed.day)}, {parsed.year}"
    return parsed.strftime(fmt)


def _qty_word(qty: object) -> str:
    try:
        number = int(qty)  # type: ignore[arg-type]
    except (TypeError, ValueError):
        return str(qty)
    return NUM_IN_WORDS.get(number, str(number))


def _expected_total_lines(hard_cases: list[str], order_data: dict) -> list[str]:
    """Riga 'Expected Total' per il caso difficile TOTAL_MIXED."""
    if "TOTAL_MIXED" not in hard_cases:
        return []
    total = (order_data.get("totals") or {}).get("total_amount_declared") or 0
    return ["", f"Expected Total: ${float(total):.2f}"]


def render_email_body(
    template: str,
    order_data: dict,
    hard_cases: list[str],
    catalog: list[dict],
) -> tuple[str, str]:
    """Costruisce il corpo di un'email. Ritorna (body, subtype plain|html).

    I cinque template sono volutamente fuori dalle regex ``EMAIL_ITEM_RE`` /
    ``EMAIL_QTY_RE`` (niente ``1. [SKU]`` né ``Qty: N UoM @ $``).
    """
    header = order_data["header"]
    buyer = order_data["buyer"]
    ship = order_data.get("ship_to") or {}
    po = header["po_number"] if "MISSING_PO_NUMBER" not in hard_cases else "[see attached]"
    date_display = header["po_date"]
    if "DATE_FORMAT" in hard_cases:
        date_display = reformat_existing_date(date_display, "eml")
    if "MISSING_DATE" in hard_cases:
        date_display = ""
    delivery = header.get("requested_delivery_date") or "ASAP"
    ship_name = ship.get("name") or ""
    lines = order_data["lines"]

    def line_bits(index: int, line: dict) -> tuple[str, object, str, str, str]:
        desc = line["description"]
        qty: object = line["quantity"]
        if "TYPO_DESC" in hard_cases and index == 0:
            desc = _add_typos(desc)
        if "AMBIGUOUS_DESC" in hard_cases and index == 1 and len(lines) > 1:
            desc = _make_ambiguous(desc, catalog)
        if "QTY_WORDS" in hard_cases and index == 0:
            qty = NUM_IN_WORDS.get(int(line["quantity"]), str(line["quantity"]))
        if "MISSING_QTY" in hard_cases and index == len(lines) - 1:
            qty = "TBD"
        code = line.get("buyer_item_code") or line.get("seller_item_code") or "N/A"
        price = f"{line['unit_price']:.2f}" if line.get("unit_price") else ""
        return code, qty, line["uom_raw"], desc, price

    if template == "classic":
        po_display = header["po_number"] if "MISSING_PO_NUMBER" not in hard_cases else "[see attached]"
        chunks = [
            "Hello,",
            "",
            "Please process the following order:",
            "",
            f"PO Number: {po_display}",
            f"Date: {date_display}",
            f"Ship To: {ship_name}",
            "",
            "Items:",
            "------",
        ]
        for index, line in enumerate(lines):
            code, qty, uom, desc, price = line_bits(index, line)
            price_str = f"@ ${price}" if price else ""
            chunks.append(f"  {line['line_no']}. [{code}] {desc}")
            chunks.append(f"     Qty: {qty} {uom} {price_str}")
        chunks.extend([
            "",
            f"Delivery requested by: {delivery}",
        ])
        chunks.extend(_expected_total_lines(hard_cases, order_data))
        chunks.extend([
            "",
            "Thank you,",
            buyer.get("contact", {}).get("name") or "Purchasing Dept",
            buyer["name"],
        ])
        return "\n".join(chunks), "plain"

    if template == "html_table":
        rows = []
        for index, line in enumerate(lines):
            code, qty, uom, desc, price = line_bits(index, line)
            rows.append(
                f"<tr><td>{code}</td><td>{desc}</td><td>{qty}</td>"
                f"<td>{uom}</td><td>{price}</td></tr>"
            )
        body = (
            "<html><body>"
            f"<p>Please process purchase order <b>{po}</b> dated {date_display}.</p>"
            f"<p>Deliver to {ship_name}. Need by {delivery}.</p>"
            "<table border='1' cellpadding='4'>"
            "<tr><th>Item</th><th>Description</th><th>Quantity</th>"
            "<th>Unit</th><th>Price</th></tr>"
            + "".join(rows)
            + "</table>"
            + (
                f"<p>Expected Total: ${float((order_data.get('totals') or {}).get('total_amount_declared') or 0):.2f}</p>"
                if "TOTAL_MIXED" in hard_cases else ""
            )
            + "<p>Thanks</p></body></html>"
        )
        return body, "html"

    if template == "free_text":
        chunks = [
            f"Hi team, could you please raise this under our purchase order {po} "
            f"dated {date_display}? Deliver to {ship_name}; we need the goods by {delivery}.",
            "",
        ]
        for index, line in enumerate(lines):
            code, qty, uom, desc, price = line_bits(index, line)
            price_bit = f" at {price} each" if price else ""
            chunks.append(
                f"We need {qty} {uom} of {desc} (catalog {code}){price_bit}."
            )
        chunks.append("")
        chunks.append("Thanks in advance.")
        chunks.extend(_expected_total_lines(hard_cases, order_data))
        return "\n".join(chunks), "plain"

    if template == "signature_list":
        chunks = [
            "Please go ahead with the replenishment below my signature.",
            "",
            "Best regards,",
            buyer.get("contact", {}).get("name") or "Purchasing",
            buyer["name"],
            "--",
            "Requested items:",
        ]
        for index, line in enumerate(lines):
            code, qty, uom, desc, price = line_bits(index, line)
            price_bit = f" / {price}" if price else ""
            chunks.append(f"{code}  |  {desc}  |  {qty} {uom}{price_bit}")
        chunks.append(f"PO {po} · {date_display} · deliver {ship_name} by {delivery}")
        chunks.extend(_expected_total_lines(hard_cases, order_data))
        return "\n".join(chunks), "plain"

    if template == "qty_words":
        chunks = [
            f"Kindly arrange shipment against purchase order {po} ({date_display}).",
            f"Ship the following to {ship_name}, needed by {delivery}:",
            "",
        ]
        for index, line in enumerate(lines):
            code, qty, uom, desc, price = line_bits(index, line)
            word = _qty_word(qty if qty != "TBD" else 0) if qty != "TBD" else "an unspecified quantity of"
            price_bit = f", unit price {price}" if price else ""
            chunks.append(
                f"- {word} {uom} of {desc}, catalog {code}{price_bit}."
            )
        chunks.append("")
        chunks.append("Appreciate the help.")
        chunks.extend(_expected_total_lines(hard_cases, order_data))
        return "\n".join(chunks), "plain"

    # forwarded_reply
    original_lines = []
    for index, line in enumerate(lines):
        code, qty, uom, desc, price = line_bits(index, line)
        price_bit = f" @ {price}" if price else ""
        original_lines.append(f"{code} x {qty} {uom}{price_bit} — {desc}")
    body = "\n".join([
        "On Mon, 12 Jan 2026 at 09:14, Warehouse <ops@example.com> wrote:",
        "> sure, please proceed with the replenishment",
        "",
        "---------- Forwarded message ---------",
        f"From: {buyer.get('contact', {}).get('name', 'Purchasing')} "
        f"<{buyer.get('contact', {}).get('email', 'orders@example.com')}>",
        "Date: Mon, 12 Jan 2026 08:02:11 +0000",
        f"Subject: Re: {po}",
        "",
        f"Could you fulfill {po} dated {date_display}?",
        f"Deliver to {ship_name} by {delivery}.",
        "Line items:",
        *original_lines,
        *_expected_total_lines(hard_cases, order_data),
    ])
    return body, "plain"


def _format_date(dt: datetime, fmt: str) -> str:
    mapping = {
        "YYYY-MM-DD": "%Y-%m-%d",
        "MM/DD/YYYY": "%m/%d/%Y",
        "DD-MMM-YYYY": "%d-%b-%Y",
        "M/D/YY": "%-m/%-d/%y",
        "YYYY/MM/DD": "%Y/%m/%d",
        "DD/MM/YYYY": "%d/%m/%Y",
        "MM-DD-YYYY": "%m-%d-%Y",
        "DD.MM.YYYY": "%d.%m.%Y",
        "MMM DD, YYYY": "%b %d, %Y",
    }
    py_fmt = mapping.get(fmt, "%Y-%m-%d")
    try:
        return dt.strftime(py_fmt)
    except ValueError:
        return dt.strftime("%Y-%m-%d")


def _add_typos(desc: str) -> str:
    typos = [
        ("Gloves", "Glovs"),
        ("Cleaner", "Clener"),
        ("Safety", "Safty"),
        ("Steel", "Stel"),
        ("industrial", "industral"),
        ("powder", "powdr"),
        ("free", "fre"),
        ("concentrate", "concetrate"),
        ("Screw", "Scew"),
        ("black", "blak"),
    ]
    for orig, typo in typos:
        if orig.lower() in desc.lower():
            return desc.replace(orig, typo, 1).replace(orig.lower(), typo.lower(), 1)
    return desc + " (approx)"


def _make_ambiguous(desc: str, catalog: list[dict]) -> str:
    desc_lower = desc.lower()
    if "liner" in desc_lower:
        return re.sub(r"\d+x\d+", "", desc)
    if "glove" in desc_lower:
        return re.sub(r"size \w+", "", desc)
    if "cleaner" in desc_lower:
        return desc.replace("concentrate", "").replace("ready to use", "")
    return desc.split()[0] + " " + " ".join(desc.split()[2:]) if len(desc.split()) > 2 else desc


def generate_pdf_order(
    order_data: dict,
    output_path: Path,
    hard_cases: list[str],
    catalog: list[dict],
    order_idx: int = 0
) -> tuple[bytes, dict]:
    """Genera un ordine in formato PDF."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )
    from reportlab.rl_config import invariant as rl_invariant

    old_invariant = rl_invariant
    import reportlab.rl_config
    reportlab.rl_config.invariant = 1

    try:
        buffer = BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter, topMargin=0.5*inch, bottomMargin=0.5*inch)
        doc.title = f"Purchase Order {order_data['header']['po_number']}"
        doc.author = "VentureOrder Synthetic Generator"
        doc.creator = "VentureOrder"
        styles = getSampleStyleSheet()
        story = []

        title_style = ParagraphStyle("Title", parent=styles["Heading1"], fontSize=16, spaceAfter=12)
        normal_style = styles["Normal"]

        story.append(Paragraph("<b>PURCHASE ORDER</b>", title_style))
        story.append(Spacer(1, 0.2*inch))

        buyer = order_data["buyer"]
        header = order_data["header"]

        po_display = header["po_number"] if "MISSING_PO_NUMBER" not in hard_cases else ""
        date_display = header["po_date"]

        if "DATE_FORMAT" in hard_cases:
            date_display = reformat_existing_date(date_display, "pdf")

        header_data = [
            ["From:", buyer["name"]],
            ["PO Number:", po_display],
            ["Date:", date_display],
            ["Ship To:", order_data.get("ship_to", {}).get("name", buyer["name"])],
        ]

        if header.get("requested_delivery_date"):
            header_data.append(["Delivery Date:", header["requested_delivery_date"]])

        header_table = Table(header_data, colWidths=[1.5*inch, 4*inch])
        header_table.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story.append(header_table)
        story.append(Spacer(1, 0.3*inch))

        lines = order_data["lines"]
        line_data = [["Line", "Item Code", "Description", "Qty", "UoM", "Price", "Amount"]]

        for i, line in enumerate(lines):
            desc = line["description"]
            qty = line["quantity"]

            if "TYPO_DESC" in hard_cases and i == 0:
                desc = _add_typos(desc)
            if "AMBIGUOUS_DESC" in hard_cases and i == 1 and len(lines) > 1:
                desc = _make_ambiguous(desc, catalog)
            if "QTY_WORDS" in hard_cases and i == 0:
                qty = NUM_IN_WORDS.get(int(qty), str(qty))
            if "MISSING_QTY" in hard_cases and i == len(lines) - 1:
                qty = ""

            line_data.append([
                str(line["line_no"]),
                line.get("buyer_item_code") or line.get("seller_item_code") or "",
                desc[:50] + "..." if len(desc) > 50 else desc,
                str(qty),
                line["uom_raw"],
                f"${line.get('unit_price', 0):.2f}" if line.get("unit_price") else "",
                f"${line.get('line_amount', 0):.2f}" if line.get("line_amount") else ""
            ])

        if "TOTAL_MIXED" in hard_cases:
            line_data.append(["", "", "*** Please confirm availability ***", "", "", "", ""])
            line_data.append(["", "", "SUBTOTAL", "", "", "", f"${order_data.get('totals', {}).get('lines_amount_declared', 0):.2f}"])

        line_table = Table(line_data, colWidths=[0.5*inch, 1*inch, 2.5*inch, 0.6*inch, 0.5*inch, 0.8*inch, 0.8*inch])
        line_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.grey),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ALIGN", (3, 1), (3, -1), "CENTER"),
            ("ALIGN", (5, 1), (-1, -1), "RIGHT"),
        ]))
        story.append(line_table)

        if "MULTIPAGE_PDF" in hard_cases:
            story.append(Spacer(1, 4*inch))
            story.append(Paragraph("<b>Page 2 - Terms and Conditions</b>", title_style))
            story.append(Spacer(1, 0.2*inch))
            for _ in range(20):
                story.append(Paragraph("Standard terms and conditions apply. Delivery within stated timeframe. "
                                       "Payment terms as agreed. All prices in USD unless otherwise specified. "
                                       "Subject to availability.", normal_style))
                story.append(Spacer(1, 0.1*inch))

        story.append(Spacer(1, 0.3*inch))
        if header.get("notes"):
            story.append(Paragraph(f"<b>Notes:</b> {header['notes']}", normal_style))

        doc.build(story)
        pdf_bytes = buffer.getvalue()
    finally:
        reportlab.rl_config.invariant = old_invariant

    doc_info = {
        "document_id": _deterministic_doc_id(order_idx, "pdf"),
        "kind": "PDF_TEXT",
        "filename": f"{header['po_number']}.pdf",
        "mime_type": "application/pdf",
        "sha256": _sha256(pdf_bytes),
        "page_count": 2 if "MULTIPAGE_PDF" in hard_cases else 1
    }

    output_path.parent.mkdir(parents=True, exist_ok=True)
    with open(output_path, "wb") as f:
        f.write(pdf_bytes)

    return pdf_bytes, doc_info


def generate_xlsx_order(
    order_data: dict,
    output_path: Path,
    hard_cases: list[str],
    catalog: list[dict],
    order_idx: int = 0
) -> tuple[bytes, dict]:
    """Genera un ordine in formato Excel."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Purchase Order"

    header_font = Font(bold=True, size=14)
    label_font = Font(bold=True)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font_white = Font(bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

    ws.merge_cells("A1:G1")
    ws["A1"] = "PURCHASE ORDER"
    ws["A1"].font = header_font
    ws["A1"].alignment = Alignment(horizontal="center")

    buyer = order_data["buyer"]
    header = order_data["header"]

    po_display = header["po_number"] if "MISSING_PO_NUMBER" not in hard_cases else ""
    date_display = header["po_date"]
    if "DATE_FORMAT" in hard_cases:
        date_display = reformat_existing_date(date_display, "xlsx")

    ws["A3"] = "Company:"
    ws["B3"] = buyer["name"]
    ws["A3"].font = label_font

    ws["A4"] = "PO Number:"
    ws["B4"] = po_display
    ws["A4"].font = label_font

    ws["A5"] = "Date:"
    ws["B5"] = date_display
    ws["A5"].font = label_font

    if "MISSING_DATE" in hard_cases:
        ws["B5"] = ""

    ws["A6"] = "Ship To:"
    ws["B6"] = order_data.get("ship_to", {}).get("name", "")
    ws["A6"].font = label_font

    ws["A7"] = "Delivery Date:"
    ws["B7"] = header.get("requested_delivery_date") or ""
    ws["A7"].font = label_font

    start_row = 9
    headers_row = ["Line", "Item Code", "Description", "Quantity", "UoM", "Unit Price", "Amount"]
    for col, h in enumerate(headers_row, 1):
        cell = ws.cell(row=start_row, column=col, value=h)
        cell.font = header_font_white
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center")

    lines = order_data["lines"]
    for i, line in enumerate(lines):
        row = start_row + 1 + i
        desc = line["description"]
        qty = line["quantity"]

        if "TYPO_DESC" in hard_cases and i == 0:
            desc = _add_typos(desc)
        if "AMBIGUOUS_DESC" in hard_cases and i == 1 and len(lines) > 1:
            desc = _make_ambiguous(desc, catalog)
        if "QTY_WORDS" in hard_cases and i == 0:
            qty = NUM_IN_WORDS.get(int(qty), str(qty))
        if "MISSING_QTY" in hard_cases and i == len(lines) - 1:
            qty = None

        ws.cell(row=row, column=1, value=line["line_no"]).border = thin_border
        ws.cell(row=row, column=2, value=line.get("buyer_item_code") or line.get("seller_item_code") or "").border = thin_border
        ws.cell(row=row, column=3, value=desc).border = thin_border
        ws.cell(row=row, column=4, value=qty).border = thin_border
        ws.cell(row=row, column=5, value=line["uom_raw"]).border = thin_border
        ws.cell(row=row, column=6, value=line.get("unit_price")).border = thin_border
        ws.cell(row=row, column=7, value=line.get("line_amount")).border = thin_border

    if "TOTAL_MIXED" in hard_cases:
        note_row = start_row + len(lines) + 1
        ws.cell(row=note_row, column=3, value="Note: Confirm stock before shipping")
        expected_row = note_row + 1
        ws.cell(row=expected_row, column=6, value="Expected Total:")
        ws.cell(row=expected_row, column=7, value=order_data.get("totals", {}).get("total_amount_declared"))
        total_row = expected_row + 1
        ws.cell(row=total_row, column=6, value="TOTAL:")
        ws.cell(row=total_row, column=7, value=order_data.get("totals", {}).get("lines_amount_declared"))

    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 15
    ws.column_dimensions["C"].width = 40
    ws.column_dimensions["D"].width = 10
    ws.column_dimensions["E"].width = 8
    ws.column_dimensions["F"].width = 12
    ws.column_dimensions["G"].width = 12

    wb.properties.created = FIXED_GENERATION_TIMESTAMP
    wb.properties.modified = FIXED_GENERATION_TIMESTAMP

    from unittest.mock import patch
    buffer = BytesIO()
    with patch('openpyxl.packaging.core.datetime') as mock_datetime:
        mock_datetime.datetime.now.return_value = FIXED_GENERATION_TIMESTAMP
        mock_datetime.datetime.utcnow.return_value = FIXED_GENERATION_TIMESTAMP
        wb.save(buffer)
    xlsx_bytes = buffer.getvalue()

    doc_info = {
        "document_id": _deterministic_doc_id(order_idx, "xlsx"),
        "kind": "XLSX",
        "filename": f"{header['po_number']}.xlsx",
        "mime_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "sha256": _sha256(xlsx_bytes),
        "page_count": None
    }

    output_path.parent.mkdir(parents=True, exist_ok=True)
    with open(output_path, "wb") as f:
        f.write(xlsx_bytes)

    return xlsx_bytes, doc_info


def generate_eml_order(
    order_data: dict,
    output_path: Path,
    hard_cases: list[str],
    catalog: list[dict],
    include_attachment: bool = False,
    order_idx: int = 0
) -> tuple[bytes, list[dict]]:
    """Genera un ordine in formato email EML."""
    buyer = order_data["buyer"]
    header = order_data["header"]
    contact = buyer.get("contact", {})

    msg = MIMEMultipart()
    msg["From"] = contact.get("email", "orders@example.com")
    msg["To"] = "orders@delta-industrial.example"
    msg["Subject"] = f"PO {header['po_number']} - Order from {buyer['name']}"
    msg["Message-ID"] = f"<po-{header['po_number']}@{buyer.get('email_domain', 'example.com')}>"
    msg["Date"] = _fixed_email_date()

    if "EMAIL_CONTRADICTORY" in hard_cases:
        # Serve il formato noto alle regex, altrimenti il preprocessor
        # non vede le righe email e prende l'allegato (quantità raddoppiate)
        # senza marcare il disaccordo.
        template = "classic"
    else:
        template = EMAIL_BODY_TEMPLATES[order_idx % len(EMAIL_BODY_TEMPLATES)]
    body_text, subtype = render_email_body(template, order_data, hard_cases, catalog)
    msg.attach(MIMEText(body_text, subtype))

    docs = []

    email_body_doc = {
        "document_id": _deterministic_doc_id(order_idx, "email-body"),
        "kind": "EMAIL_BODY",
        "filename": None,
        "mime_type": "text/plain",
        "sha256": _sha256(body_text.encode("utf-8")),
        "page_count": None
    }
    docs.append(email_body_doc)

    if include_attachment or "EMAIL_CONTRADICTORY" in hard_cases:
        if "EMAIL_CONTRADICTORY" in hard_cases:
            modified_order = json.loads(json.dumps(order_data))
            for line in modified_order["lines"]:
                line["quantity"] = line["quantity"] * 2
                if line.get("line_amount"):
                    line["line_amount"] = line["line_amount"] * 2
            attach_bytes, attach_doc = generate_xlsx_order(
                modified_order,
                output_path.with_suffix(".xlsx"),
                [],
                catalog,
                order_idx=order_idx
            )
        else:
            attach_bytes, attach_doc = generate_xlsx_order(
                order_data,
                output_path.with_suffix(".xlsx"),
                [],
                catalog,
                order_idx=order_idx
            )

        attachment = MIMEApplication(attach_bytes, _subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        attachment.add_header("Content-Disposition", "attachment", filename=f"PO_{header['po_number']}.xlsx")
        msg.attach(attachment)
        docs.append(attach_doc)

    eml_bytes = msg.as_bytes()

    output_path.parent.mkdir(parents=True, exist_ok=True)
    with open(output_path, "wb") as f:
        f.write(eml_bytes)

    return eml_bytes, docs


def _create_ground_truth(
    order_data: dict,
    documents: list[dict],
    channel: str,
    hard_cases: list[str],
    order_idx: int = 0
) -> dict:
    """Crea il ground truth JSON conforme allo schema VOF."""
    order_id = _deterministic_uuid(order_idx, "order")
    buyer = order_data["buyer"]
    header = order_data["header"]

    status = "VALIDATED"
    validation_issues = []

    if "MISSING_PO_NUMBER" in hard_cases:
        status = "NEEDS_REVIEW"
        validation_issues.append({
            "code": "HDR-01",
            "severity": "BLOCK",
            "path": "/header/po_number",
            "message": "Manca il numero d'ordine (PO number)"
        })

    if "MISSING_QTY" in hard_cases:
        status = "NEEDS_REVIEW"
        validation_issues.append({
            "code": "QTY-01",
            "severity": "BLOCK",
            "path": f"/lines/{len(order_data['lines'])-1}/quantity",
            "message": "Quantità mancante o non valida"
        })

    if "MISSING_DATE" in hard_cases:
        status = "NEEDS_REVIEW"
        validation_issues.append({
            "code": "HDR-02",
            "severity": "REVIEW",
            "path": "/header/po_date",
            "message": "Data ordine mancante o non plausibile"
        })

    if "EMAIL_CONTRADICTORY" in hard_cases:
        status = "NEEDS_REVIEW"
        validation_issues.append({
            "code": "ITM-04",
            "severity": "REVIEW",
            "path": "/lines",
            "message": "Quantità nel corpo email diverse dall'allegato"
        })

    if "AMBIGUOUS_DESC" in hard_cases:
        status = "NEEDS_REVIEW"
        validation_issues.append({
            "code": "ITM-03",
            "severity": "REVIEW",
            "path": "/lines/1/description",
            "message": "Descrizione ambigua, richiede conferma articolo"
        })

    if "UOM_CONVERSION" in hard_cases or "UOM_CARTON" in hard_cases:
        validation_issues.append({
            "code": "UOM-03",
            "severity": "INFO",
            "path": "/lines/0",
            "message": "Conversione da cartone applicata"
        })

    if "UNKNOWN_CUSTOMER_CODE" in hard_cases:
        status = "NEEDS_REVIEW"
        validation_issues.append({
            "code": "ITM-01",
            "severity": "BLOCK",
            "path": "/lines/0/resolution",
            "message": "Codice articolo cliente sconosciuto: riga non risolta"
        })

    if "AMBIGUOUS_CUSTOMER_CODE" in hard_cases:
        status = "NEEDS_REVIEW"
        validation_issues.append({
            "code": "ITM-04",
            "severity": "REVIEW",
            "path": "/lines/0",
            "message": "Codice articolo cliente ambiguo: identificatori in conflitto"
        })

    if "PRICE_OUT_OF_TOLERANCE" in hard_cases:
        status = "NEEDS_REVIEW"
        validation_issues.append({
            "code": "PRC-01",
            "severity": "REVIEW",
            "path": "/lines/0/unit_price",
            "message": "Prezzo PO fuori tolleranza rispetto al listino"
        })

    if "DUP_PO" in hard_cases:
        status = "NEEDS_REVIEW"
        validation_issues.append({
            "code": "DUP-01",
            "severity": "BLOCK",
            "path": "/header/po_number",
            "message": "PO duplicato: stesso cliente, numero e revisione già POSTED"
        })

    lines_gt = []
    for line in order_data["lines"]:
        line_gt = {
            "line_no": line["line_no"],
            "buyer_item_code": line.get("buyer_item_code"),
            "seller_item_code": line.get("seller_item_code"),
            "manufacturer_name": line.get("manufacturer_name"),
            "manufacturer_part_number": line.get("manufacturer_part_number"),
            "gtin": line.get("gtin"),
            "description": line["description"],
            "quantity": line["quantity"],
            "uom_raw": line["uom_raw"],
            "uom": line.get("uom"),
            "pack": line.get("pack"),
            "unit_price": line.get("unit_price"),
            "price_uom": line.get("price_uom"),
            "price_basis_qty": line.get("price_basis_qty"),
            "discount_pct": line.get("discount_pct"),
            "line_amount": line.get("line_amount"),
            "requested_delivery_date": line.get("requested_delivery_date"),
            "notes": line.get("notes"),
            "resolution": line.get("resolution")
        }
        lines_gt.append(line_gt)

    ground_truth = {
        "schema_version": "1.0",
        "order_id": order_id,
        "tenant_id": "delta-industrial-supply",
        "document_type": "ORIGINAL",
        "status": status,
        "source": {
            "channel": channel,
            "message_id": f"<po-{header['po_number']}@{buyer.get('email_domain', 'example.com')}>" if channel.startswith("EMAIL") else None,
            "from_address": buyer.get("contact", {}).get("email") if channel.startswith("EMAIL") else None,
            "subject": f"PO {header['po_number']} - Order from {buyer['name']}" if channel.startswith("EMAIL") else None,
            "received_at": _fixed_timestamp_iso(),
            "documents": documents
        },
        "buyer": {
            "name": buyer["name"],
            "code_on_document": buyer.get("erp_partner_ref"),
            "tax_id": buyer.get("tax_id"),
            "gln": None,
            "peppol_endpoint": None,
            "address": buyer.get("address"),
            "contact": buyer.get("contact"),
            "resolved": {
                "erp_partner_id": buyer.get("erp_partner_id"),
                "erp_partner_ref": buyer.get("erp_partner_ref"),
                "method": "EMAIL_DOMAIN" if channel.startswith("EMAIL") else "CODE",
                "score": 0.99
            }
        },
        "seller": {
            "name": "Delta Industrial Supply LLC (SYNTHETIC)",
            "code_on_document": None,
            "tax_id": None,
            "address": {
                "line1": "55 Commerce Dr",
                "city": "Dayton",
                "region": "OH",
                "postal_code": "45402",
                "country": "US"
            },
            "resolved": None
        },
        "header": header,
        "lines": lines_gt,
        "charges": [],
        "totals": order_data.get("totals"),
        "decision": {
            "score": 0.75 if validation_issues else 0.98,
            "route": "REVIEW" if status == "NEEDS_REVIEW" else "AUTO_POST",
            "reasons": [i["code"] for i in validation_issues] if validation_issues else []
        },
        "processing": {
            "extractor": "synthetic/generator.py",
            "prompt_version": None,
            "created_at": _fixed_timestamp_iso(),
            "updated_at": None
        }
    }

    if order_data.get("ship_to"):
        ground_truth["ship_to"] = order_data["ship_to"]

    if validation_issues:
        ground_truth["validation"] = {"issues": validation_issues}

    if hard_cases:
        ground_truth["_synthetic_metadata"] = {
            "hard_case_types": hard_cases,
            "is_hard_case": True
        }
    else:
        ground_truth["_synthetic_metadata"] = {
            "hard_case_types": [],
            "is_hard_case": False
        }

    return ground_truth


def generate_orders(
    catalog: list[dict],
    customers: list[dict],
    output_dir: Path,
    num_orders: int = 150,
    hard_case_pct: float = 0.15
) -> tuple[list[dict], dict]:
    """Genera tutti gli ordini con documenti e ground truth.

    Args:
        catalog: Lista articoli.
        customers: Lista clienti.
        output_dir: Directory di output.
        num_orders: Numero totale ordini.
        hard_case_pct: Percentuale casi difficili.

    Returns:
        (lista ground truth, manifest)
    """
    random.seed(SEED + 2)

    num_hard = int(num_orders * hard_case_pct)
    order_configs = assign_hard_cases(num_orders, hard_case_pct)

    pdf_dir = output_dir / "documents" / "pdf"
    xlsx_dir = output_dir / "documents" / "xlsx"
    eml_dir = output_dir / "documents" / "eml"
    gt_dir = output_dir / "ground_truth"

    for folder in [pdf_dir, xlsx_dir, eml_dir, gt_dir]:
        folder.mkdir(parents=True, exist_ok=True)
        for old in folder.iterdir():
            if old.is_file():
                old.unlink()

    ground_truths = []
    manifest_entries = []

    base_date = datetime(2026, 1, 15)
    po_counter = 1000

    formats_count = {"PDF": 0, "XLSX": 0, "EML": 0}
    target_per_format = num_orders // 3
    last_po_by_customer: dict[str, str] = {}

    for order_idx, hard_cases in enumerate(order_configs):
        customer = random.choice(customers)
        if "DUP_PO" in hard_cases and last_po_by_customer:
            prior_cust = random.choice(list(last_po_by_customer))
            customer = next(c for c in customers if c["customer_id"] == prior_cust)
        style = customer["order_style"]

        available_formats = []
        for fmt in style["format_preference"]:
            if formats_count[fmt] < target_per_format + 5:
                available_formats.append(fmt)
        if not available_formats:
            available_formats = ["PDF", "XLSX", "EML"]

        if "MULTIPAGE_PDF" in hard_cases:
            doc_format = "PDF"
        elif "EMAIL_CONTRADICTORY" in hard_cases:
            doc_format = "EML"
        else:
            doc_format = random.choice(available_formats)

        formats_count[doc_format] += 1

        po_number = f"{style['po_prefix']}{po_counter:05d}"
        po_counter += 1
        if "DUP_PO" in hard_cases and customer["customer_id"] in last_po_by_customer:
            po_number = last_po_by_customer[customer["customer_id"]]

        order_date = base_date + timedelta(days=random.randint(0, 180))
        delivery_date = order_date + timedelta(days=random.randint(5, 21))

        num_lines = random.randint(*style["typical_lines"])
        selected_items = random.sample(catalog, min(num_lines, len(catalog)))
        if "UOM_CARTON" in hard_cases or "UOM_CONVERSION" in hard_cases:
            carton_items = [entry for entry in catalog if carton_conversion(entry)]
            if carton_items:
                first = random.choice(carton_items)
                rest = [entry for entry in selected_items if entry["sku"] != first["sku"]]
                selected_items = [first] + rest[: max(0, len(selected_items) - 1)]

        lines = []
        total_amount = 0.0

        for line_no, item in enumerate(selected_items, 1):
            qty_base = random.choice([1, 2, 3, 4, 5, 6, 10, 12, 24])
            line_uom = item["uom"]
            erp_uom = item["uom"]
            conversion_factor = 1.0
            erp_quantity = qty_base
            gtin = None
            mpn = None

            if ("UOM_CARTON" in hard_cases or "UOM_CONVERSION" in hard_cases) and line_no == 1:
                conv = carton_conversion(item)
                if conv:
                    uom_raw = conv["from_uom"]
                    line_uom = "CS"
                    erp_uom = conv["to_uom"]
                    conversion_factor = float(conv["factor"])
                    erp_quantity = qty_base * conversion_factor
                elif item.get("pack"):
                    pack = item["pack"]
                    uom_raw = f"CS of {pack['inner_count']}"
                else:
                    uom_raw = random.choice(["DZ", "DZN", "DOZEN"])
                    qty_base = random.choice([1, 2, 3])
            else:
                uom_raw = item["uom"]

            buyer_item_code = None
            seller_item_code = None

            if style.get("uses_own_codes") and customer.get("xref"):
                xref_match = next((x for x in customer["xref"] if x["sku"] == item["sku"]), None)
                if xref_match:
                    if "CUSTOMER_CODES" in hard_cases or random.random() < 0.7:
                        buyer_item_code = xref_match["customer_code"]
                    else:
                        seller_item_code = item["sku"]
                else:
                    seller_item_code = item["sku"]
            else:
                seller_item_code = item["sku"]

            if "UNKNOWN_CUSTOMER_CODE" in hard_cases and line_no == 1:
                prefix = style.get("code_prefix") or customer["customer_id"].replace("C-", "") + "-"
                buyer_item_code = f"{prefix}UNKNOWN-{order_idx:04d}"
                seller_item_code = None

            if "AMBIGUOUS_CUSTOMER_CODE" in hard_cases and line_no == 1:
                other = next((entry for entry in catalog if entry["sku"] != item["sku"]), None)
                if other and other.get("gtin"):
                    gtin = other["gtin"]
                    mpn = other.get("manufacturer_part_number")
                if not buyer_item_code:
                    xref_any = next(iter(customer.get("xref") or []), None)
                    if xref_any:
                        buyer_item_code = xref_any["customer_code"]
                    else:
                        buyer_item_code = f"{customer['customer_id']}-AMB"

            unit_price = item["price"]
            list_price = item["price"]
            if "PRICE_OUT_OF_TOLERANCE" in hard_cases and line_no == 1:
                unit_price = round(float(item["price"]) * 1.15, 2)
            line_amount = round(qty_base * unit_price, 2)
            total_amount += line_amount

            unresolved = "UNKNOWN_CUSTOMER_CODE" in hard_cases and line_no == 1
            ambiguous = "AMBIGUOUS_CUSTOMER_CODE" in hard_cases and line_no == 1
            if unresolved:
                method, score = "UNRESOLVED", 0.0
            elif ambiguous:
                method, score = "FUZZY", 0.82
            elif buyer_item_code:
                method, score = "XREF", 1.0
            else:
                method, score = "SELLER_CODE", 0.98
            candidates = []
            if ambiguous:
                other_sku = item["sku"]
                if gtin:
                    other_sku = next(
                        (entry["sku"] for entry in catalog if entry.get("gtin") == gtin),
                        item["sku"],
                    )
                candidates = [
                    {"sku": item["sku"], "score": 0.82},
                    {"sku": other_sku, "score": 0.80},
                ]
            resolution = {
                "sku": None if unresolved else item["sku"],
                "erp_product_id": None if unresolved else item["erp_product_id"],
                "erp_uom": erp_uom,
                "erp_uom_id": random.randint(1, 50),
                "erp_quantity": erp_quantity,
                "conversion_factor": conversion_factor,
                "method": method,
                "score": score,
                "candidates": candidates,
                "erp_list_price": list_price,
            }

            line = {
                "line_no": line_no,
                "buyer_item_code": buyer_item_code,
                "seller_item_code": seller_item_code,
                "manufacturer_name": None,
                "manufacturer_part_number": mpn,
                "gtin": gtin,
                "description": item["description"],
                "quantity": qty_base,
                "uom_raw": uom_raw,
                "uom": line_uom,
                "pack": item.get("pack"),
                "unit_price": unit_price,
                "price_uom": line_uom,
                "price_basis_qty": 1,
                "discount_pct": None,
                "line_amount": line_amount,
                "requested_delivery_date": None,
                "notes": None,
                "resolution": resolution,
            }
            lines.append(line)

        last_po_by_customer[customer["customer_id"]] = po_number

        ship_to_choice = random.choice(customer["ship_to"])
        ship_to = {
            "name": ship_to_choice["name"],
            "code_on_document": ship_to_choice["code"],
            "address": ship_to_choice["address"],
            "resolved": {
                "erp_partner_id": customer["erp_partner_id"] + 1,
                "erp_partner_ref": ship_to_choice["code"],
                "method": "CODE",
                "score": 1.0
            }
        }

        order_data = {
            "buyer": {
                "name": customer["name"],
                "tax_id": customer.get("tax_id"),
                "erp_partner_id": customer["erp_partner_id"],
                "erp_partner_ref": customer["erp_partner_ref"],
                "email_domain": customer.get("email_domain"),
                "contact": customer.get("contact"),
                "address": customer.get("address")
            },
            "ship_to": ship_to,
            "header": {
                "po_number": po_number,
                "po_revision": 0,
                "po_date": _format_date(order_date, style["date_format"]),
                "requested_delivery_date": _format_date(delivery_date, "YYYY-MM-DD"),
                "currency": style["currency"],
                "buyer_reference": f"Dept {random.randint(100, 999)}",
                "payment_terms_text": random.choice(["Net 30", "Net 45", "2% 10 Net 30"]),
                "incoterms": None,
                "ship_via": random.choice(["Best Way", "UPS Ground", "FedEx", None]),
                "notes": random.choice([None, "Rush order", "Confirm availability", "Call before delivery"]),
                "language": style["language"]
            },
            "lines": lines,
            "totals": {
                "lines_amount_declared": round(total_amount, 2),
                "tax_amount_declared": None,
                "total_amount_declared": round(total_amount, 2),
                "lines_amount_computed": round(total_amount, 2)
            }
        }

        filename_base = f"order_{order_idx+1:03d}_{po_number.replace('/', '-')}"

        if doc_format == "PDF":
            channel = "FOLDER"
            _, doc_info = generate_pdf_order(
                order_data,
                pdf_dir / f"{filename_base}.pdf",
                hard_cases,
                catalog,
                order_idx=order_idx
            )
            documents = [doc_info]
        elif doc_format == "XLSX":
            channel = "FOLDER"
            _, doc_info = generate_xlsx_order(
                order_data,
                xlsx_dir / f"{filename_base}.xlsx",
                hard_cases,
                catalog,
                order_idx=order_idx
            )
            documents = [doc_info]
        else:
            channel = "EMAIL_IMAP"
            include_attach = random.random() < 0.3 or "EMAIL_CONTRADICTORY" in hard_cases
            _, documents = generate_eml_order(
                order_data,
                eml_dir / f"{filename_base}.eml",
                hard_cases,
                catalog,
                include_attachment=include_attach,
                order_idx=order_idx
            )

        ground_truth = _create_ground_truth(order_data, documents, channel, hard_cases, order_idx=order_idx)

        gt_path = gt_dir / f"{filename_base}.json"
        with open(gt_path, "w", encoding="utf-8") as f:
            json.dump(ground_truth, f, indent=2, ensure_ascii=False)

        ground_truths.append(ground_truth)

        manifest_entries.append({
            "order_index": order_idx + 1,
            "po_number": po_number,
            "customer_id": customer["customer_id"],
            "customer_name": customer["name"],
            "format": doc_format,
            "num_lines": len(lines),
            "total_amount": round(total_amount, 2),
            "is_hard_case": bool(hard_cases),
            "hard_case_types": hard_cases,
            "document_path": str(gt_path.relative_to(output_dir)),
            "ground_truth_path": str(gt_path.relative_to(output_dir))
        })

    manifest = {
        "generated_at": _fixed_timestamp_iso(),
        "seed": SEED,
        "total_orders": num_orders,
        "hard_cases": num_hard,
        "hard_case_percentage": round(num_hard / num_orders * 100, 1),
        "format_distribution": formats_count,
        "orders": manifest_entries
    }

    manifest_path = output_dir / "data" / "manifest.json"
    manifest_path.parent.mkdir(parents=True, exist_ok=True)
    with open(manifest_path, "w", encoding="utf-8") as f:
        json.dump(manifest, f, indent=2, ensure_ascii=False)

    import csv
    csv_path = output_dir / "data" / "manifest.csv"
    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=[
            "order_index", "po_number", "customer_id", "customer_name",
            "format", "num_lines", "total_amount", "is_hard_case",
            "hard_case_types", "document_path", "ground_truth_path"
        ])
        writer.writeheader()
        for entry in manifest_entries:
            row = entry.copy()
            row["hard_case_types"] = ";".join(entry["hard_case_types"])
            writer.writerow(row)

    return ground_truths, manifest


if __name__ == "__main__":
    from .catalog import generate_catalog
    from .customers import generate_customers

    output_dir = Path(__file__).parent

    catalog = generate_catalog(output_dir / "data" / "catalog.json")
    print(f"Generated {len(catalog)} catalog items")

    customers = generate_customers(catalog, output_dir / "data" / "customers.json")
    print(f"Generated {len(customers)} customers")

    ground_truths, manifest = generate_orders(catalog, customers, output_dir)
    print(f"Generated {len(ground_truths)} orders")
    print(f"Hard cases: {manifest['hard_cases']} ({manifest['hard_case_percentage']}%)")
    print(f"Format distribution: {manifest['format_distribution']}")
