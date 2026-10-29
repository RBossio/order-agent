"""Parser per file Excel (.xlsx, .xls)."""

from __future__ import annotations

import hashlib
import uuid
from io import BytesIO
from pathlib import Path
from typing import Any

from ..models import (
    Cell,
    FileType,
    NormalizedInput,
    Part,
    PartType,
    Source,
    Table,
    Warning,
)


def _get_cell_ref(row: int, col: int) -> str:
    """Converte indici riga/colonna in riferimento cella Excel (es. A1, B2)."""
    result = ""
    col_num = col + 1
    while col_num > 0:
        col_num, remainder = divmod(col_num - 1, 26)
        result = chr(65 + remainder) + result
    return f"{result}{row + 1}"


class ExcelParser:
    """Parser per file Excel (.xlsx, .xls)."""

    def __init__(self):
        pass

    def parse(self, file_path: Path | str) -> NormalizedInput:
        """Parsa un file Excel."""
        file_path = Path(file_path)
        
        with open(file_path, "rb") as f:
            content = f.read()
        
        file_ext = file_path.suffix.lower()
        file_type = FileType.XLSX if file_ext == ".xlsx" else FileType.XLS
        
        return self.parse_bytes(
            content,
            filename=file_path.name,
            file_path=str(file_path),
            file_type=file_type
        )

    def parse_bytes(
        self,
        content: bytes,
        filename: str = "document.xlsx",
        file_path: str | None = None,
        file_type: FileType = FileType.XLSX
    ) -> NormalizedInput:
        """Parsa contenuto Excel da bytes."""
        sha256 = hashlib.sha256(content).hexdigest()
        source = Source(
            filename=filename,
            sha256=sha256,
            file_type=file_type,
            file_path=file_path,
            file_size_bytes=len(content),
        )
        
        result = NormalizedInput(source=source)
        
        if file_type == FileType.XLS:
            self._parse_xls(content, result)
        else:
            self._parse_xlsx(content, result)
        
        return result

    def _parse_xlsx(self, content: bytes, result: NormalizedInput) -> None:
        """Parsa file .xlsx usando openpyxl."""
        from openpyxl import load_workbook
        from openpyxl.cell import MergedCell
        
        buffer = BytesIO(content)
        
        try:
            wb = load_workbook(buffer, data_only=True, read_only=False)
            
            for sheet_idx, sheet_name in enumerate(wb.sheetnames):
                ws = wb[sheet_name]
                part = self._process_xlsx_sheet(ws, sheet_name, sheet_idx, result)
                result.parts.append(part)
                
        except Exception as e:
            result.warnings.append(Warning(
                code="XLS-01",
                message=f"Errore nel parsing del file Excel: {str(e)}",
                details={"error": str(e)},
            ))

    def _process_xlsx_sheet(
        self,
        ws: Any,
        sheet_name: str,
        sheet_idx: int,
        result: NormalizedInput
    ) -> Part:
        """Processa un singolo foglio .xlsx."""
        part = Part(
            part_type=PartType.EXCEL_SHEET,
            sheet_name=sheet_name,
            sheet_index=sheet_idx,
            mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        
        merged_ranges: dict[str, str] = {}
        if hasattr(ws, "merged_cells") and ws.merged_cells:
            for merge_range in ws.merged_cells.ranges:
                for row in range(merge_range.min_row, merge_range.max_row + 1):
                    for col in range(merge_range.min_col, merge_range.max_col + 1):
                        cell_ref = _get_cell_ref(row - 1, col - 1)
                        merged_ranges[cell_ref] = str(merge_range)
        
        table = Table()
        text_lines = []
        
        max_row = ws.max_row or 0
        max_col = ws.max_column or 0
        
        first_data_row = None
        last_data_row = None
        
        for row_idx in range(1, max_row + 1):
            row_has_data = False
            for col_idx in range(1, max_col + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                if cell.value is not None:
                    row_has_data = True
                    break
            
            if row_has_data:
                if first_data_row is None:
                    first_data_row = row_idx
                last_data_row = row_idx
        
        if first_data_row is None:
            part.text = ""
            part.compute_sha256("")
            return part
        
        for row_idx in range(first_data_row, last_data_row + 1):
            row_cells: list[Cell] = []
            row_text_parts: list[str] = []
            
            for col_idx in range(1, max_col + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell_ref = _get_cell_ref(row_idx - 1, col_idx - 1)
                
                value = cell.value
                if value is not None and not isinstance(value, (str, int, float, bool)):
                    value = str(value)
                
                is_merged = cell_ref in merged_ranges
                merge_range = merged_ranges.get(cell_ref)
                
                table_cell = Cell(
                    value=value,
                    raw_value=str(value) if value is not None else None,
                    cell_ref=cell_ref,
                    is_merged=is_merged,
                    merge_range=merge_range,
                )
                row_cells.append(table_cell)
                
                if value is not None:
                    row_text_parts.append(str(value))
            
            table.rows.append(row_cells)
            if row_text_parts:
                text_lines.append("\t".join(row_text_parts))
        
        header_idx = self._detect_header_row(table.rows)
        table.header_row_index = header_idx
        part.tables = [table]
        
        part.text = "\n".join(text_lines)
        part.compute_sha256(part.text)
        
        return part

    def _parse_xls(self, content: bytes, result: NormalizedInput) -> None:
        """Parsa file .xls usando xlrd."""
        try:
            import xlrd
        except ImportError:
            result.warnings.append(Warning(
                code="XLS-02",
                message="Libreria xlrd non disponibile per file .xls",
                details={"suggestion": "Installa xlrd con: pip install xlrd"},
            ))
            return
        
        buffer = BytesIO(content)
        
        try:
            wb = xlrd.open_workbook(file_contents=content)
            
            for sheet_idx in range(wb.nsheets):
                ws = wb.sheet_by_index(sheet_idx)
                part = self._process_xls_sheet(ws, sheet_idx, result)
                result.parts.append(part)
                
        except Exception as e:
            result.warnings.append(Warning(
                code="XLS-03",
                message=f"Errore nel parsing del file .xls: {str(e)}",
                details={"error": str(e)},
            ))

    def _process_xls_sheet(
        self,
        ws: Any,
        sheet_idx: int,
        result: NormalizedInput
    ) -> Part:
        """Processa un singolo foglio .xls."""
        part = Part(
            part_type=PartType.EXCEL_SHEET,
            sheet_name=ws.name,
            sheet_index=sheet_idx,
            mime_type="application/vnd.ms-excel",
        )
        
        table = Table()
        text_lines = []
        
        for row_idx in range(ws.nrows):
            row_cells: list[Cell] = []
            row_text_parts: list[str] = []
            
            for col_idx in range(ws.ncols):
                cell = ws.cell(row_idx, col_idx)
                value = cell.value
                cell_ref = _get_cell_ref(row_idx, col_idx)
                
                if value == "":
                    value = None
                
                table_cell = Cell(
                    value=value,
                    raw_value=str(value) if value is not None else None,
                    cell_ref=cell_ref,
                )
                row_cells.append(table_cell)
                
                if value is not None:
                    row_text_parts.append(str(value))
            
            table.rows.append(row_cells)
            if row_text_parts:
                text_lines.append("\t".join(row_text_parts))
        
        header_idx = self._detect_header_row(table.rows)
        table.header_row_index = header_idx
        part.tables = [table]
        
        part.text = "\n".join(text_lines)
        part.compute_sha256(part.text)
        
        return part

    def _detect_header_row(self, rows: list[list[Cell]]) -> int | None:
        """Rileva euristicamente la riga di intestazione."""
        if not rows:
            return None
        
        header_keywords = {
            "item", "qty", "quantity", "price", "amount", "description",
            "code", "sku", "unit", "uom", "line", "total", "product",
            "articolo", "quantità", "prezzo", "importo", "descrizione",
            "codice", "unità", "riga", "totale", "prodotto",
            "company", "po number", "date", "ship to",
        }
        
        for idx, row in enumerate(rows[:10]):
            text_cells = 0
            header_matches = 0
            non_empty = 0
            
            for cell in row:
                if cell.value is not None:
                    non_empty += 1
                    if isinstance(cell.value, str):
                        text_cells += 1
                        cell_lower = cell.value.lower().strip()
                        if cell_lower in header_keywords:
                            header_matches += 1
                        for kw in header_keywords:
                            if kw in cell_lower:
                                header_matches += 0.5
                                break
            
            if header_matches >= 2:
                return idx
        
        for idx, row in enumerate(rows[:10]):
            non_empty = sum(1 for c in row if c.value is not None)
            if non_empty >= 3:
                all_text = all(
                    isinstance(c.value, str) for c in row if c.value is not None
                )
                if all_text:
                    next_row_has_nums = False
                    if idx + 1 < len(rows):
                        for c in rows[idx + 1]:
                            if isinstance(c.value, (int, float)):
                                next_row_has_nums = True
                                break
                    if next_row_has_nums:
                        return idx
        
        return None
